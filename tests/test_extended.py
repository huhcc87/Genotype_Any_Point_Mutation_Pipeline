from io import BytesIO
from docx import Document
from src.document_input import extract_sequences
from src.ranking import rank_enzymes
from src.ladders import LADDERS
from src.visualization import sequence_markup

def test_plain_text_sequence_extraction():
    sequence="ACGT"*40
    items,warnings=extract_sequences("input.txt",("Sequence:\n"+sequence).encode())
    assert items[0]["sequence"]==sequence

def test_multiple_fasta_records():
    data=">first\n"+("ACGT"*20)+"\n>second\n"+("TGCA"*20)+"\n"
    items,warnings=extract_sequences("sample.fa",data.encode())
    assert len(items)==2
    assert len(warnings)>=1

def test_document_extraction():
    doc=Document()
    doc.add_paragraph("Example sequence")
    doc.add_paragraph("ATGC"*20)
    buf=BytesIO()
    doc.save(buf)
    records,warnings=extract_sequences("dna.docx",buf.getvalue())
    assert records[0]["length"]==80

def test_invalid_file():
    import pytest
    with pytest.raises(ValueError):
        extract_sequences("bad.xls",b"AGCT"*30)

def test_rank_enzymes():
    base={
        "enzyme":"TestI","fragments_wt":[200],
        "fragments_mut":[120,80],"cuts_wt":[],
        "cuts_mut":[120],"preferred":False,
    }
    grades=rank_enzymes([base])
    assert grades[0]["priority_rank"]==1
    assert grades[0]["design_grade"] in ("A","B","C","D")

def test_ladders():
    assert len(LADDERS)>=4
    for entry in LADDERS.values():
        assert all(isinstance(bp,int) and bp>0 for bp in entry["sizes"])
