"""Extract candidate DNA strings from user-uploaded files without guessing coordinates."""
from io import BytesIO
from pathlib import Path
import re
from .sequence import load_sequence

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MIN_CANDIDATE_BP = 60
FORMATS = [".fasta", ".fa", ".fna", ".txt", ".docx", ".pdf"]


def _extract_text(filename, data):
    suffix = Path(filename).suffix.lower()
    if suffix not in FORMATS:
        raise ValueError("Supported formats: FASTA (.fa/.fasta/.fna), TXT, DOCX, or text-based PDF.")
    if not isinstance(data, bytes) or not data or len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File must be nonempty and at most 10 MB.")
    if suffix == ".docx":
        from docx import Document
        document = Document(BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.append(cell.text)
        return "\n".join(lines)
    if suffix == ".pdf":
        from pypdf import PdfReader
        pdf = PdfReader(BytesIO(data), strict=False)
        if len(pdf.pages) > 200:
            raise ValueError("PDF exceeds 200 pages; please extract the relevant DNA sequence.")
        content = "\n".join((page.extract_text() or "") for page in pdf.pages)
        if not content.strip():
            raise ValueError("No selectable text found in PDF. Scanned documents require OCR before upload.")
        return content
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Text sequence file must use UTF-8 encoding.") from exc


def extract_sequences(filename, data):
    """Return (candidate dictionaries, warnings). Never silently choose among sequences."""
    text = _extract_text(filename, data)
    warnings = []
    suffix = Path(filename).suffix.lower()
    candidates = []

    def add(label, raw):
        try:
            sequence = load_sequence(raw)
        except ValueError:
            return
        if len(sequence) < MIN_CANDIDATE_BP:
            return
        if not any(item["sequence"] == sequence for item in candidates):
            candidates.append({"name": label[:100], "sequence": sequence, "length": len(sequence)})

    if suffix in (".fa", ".fasta", ".fna"):
        if not text.lstrip().startswith(">"):
            raise ValueError("FASTA upload must start with a >header and DNA lines.")
    if text.lstrip().startswith(">"):
        records = re.split(r"(?=^>)", text, flags=re.MULTILINE)
        for record in records:
            if not record.lstrip().startswith(">"):
                continue
            lines = record.splitlines()
            header = lines[0][1:].strip() or "FASTA record"
            add(header, record)
    if not candidates:
        block = []
        index = 0

        def flush():
            nonlocal index
            if block:
                index += 1
                add(f"DNA block {index}", "".join(block))
                block.clear()

        for line in text.splitlines():
            clean = re.sub(r"\s+", "", line)
            if re.fullmatch(r"[ACGTacgt]+", clean or " ") and len(clean) >= 12:
                block.append(clean)
            else:
                flush()
                match = re.search(r"(?:DNA\s*sequence|sequence|seq)\s*[:=]\s*([ACGTacgt]{60,})", line, re.I)
                if match:
                    add(f"Labelled sequence {index + 1}", match.group(1))
        flush()
    if not candidates:
        try:
            add("Complete file", text)
        except ValueError:
            pass
    if not candidates:
        raise ValueError(
            "No continuous A/C/G/T sequence of at least 60 bp found. "
            "For DOCX/PDF put DNA on its own lines; scanned PDFs require OCR first."
        )
    if len(candidates) > 1:
        warnings.append("Multiple sequences found. Select the intended reference sequence; coordinates refer to that sequence only.")
    if suffix in (".docx", ".pdf"):
        warnings.append("Sequence extracted from document text; verify the DNA bases against your reference before using results.")
    return candidates, warnings
