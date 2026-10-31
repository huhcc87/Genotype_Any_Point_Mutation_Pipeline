"""One-click reproducible exports: PDF, DOCX, HTML, CSV, JSON, FASTA, PNG, ZIP."""
import io
import json
from html import escape
from pathlib import Path
import tempfile
import zipfile

import pandas as pd
from docx import Document
from docx.shared import Inches
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as PDFImage
)
from .ladders import LADDERS, ladder_notes
from .visualization import render_ladder_gel, render_sequence_figure, sequence_markup

DISCLAIMER = (
    "IN-SILICO RESEARCH USE ONLY. Predictions are not experimental "
    "confirmation, clinical interpretation or diagnostic genotype calls. "
    "Verify input coordinates, genomic specificity, enzyme properties, and gel bands."
)


def _summary_rows(result, enzyme, ladder_name):
    pair=result["selected_primer"]
    rows=[
        ("Assay name",result["name"]),("Variant",result["mutation"]),
        ("Coordinates",result["coordinate_system"]),
        ("Reference length",f"{result['seq_length']} bp"),
        ("Status",result["status"]),
        ("Selected enzyme",enzyme["enzyme"] if enzyme else "None"),
        ("Enzyme design grade",enzyme.get("design_grade","—") if enzyme else "—"),
        ("Score (heuristic)",str(enzyme.get("priority_score","—")) if enzyme else "—"),
        ("Selected ladder",ladder_name),
    ]
    meta=result.get("variant_input") or {}
    if meta.get("reference_codon"):
        rows += [("Transcript",meta.get("transcript_id","")),("Genome assembly",meta.get("assembly","")),("Genomic position",str(meta.get("genomic_chromosome",""))+":"+str(meta.get("genomic_position",""))),("Coding codon",str(meta.get("codon_number",""))),("WT to mutant codon",str(meta["reference_codon"])+" → "+str(meta["alternate_codon"])),("Predicted amino acid change",str(meta.get("reference_aa",""))+str(meta.get("codon_number",""))+str(meta.get("alternate_aa","")))]
    if pair:
        rows += [
            ("Forward (5' to 3')",pair["forward"]),
            ("Reverse (5' to 3')",pair["reverse"]),
            ("Amplicon size",f"{pair['product_size']} bp"),
            ("Forward Tm",f"{pair['forward_tm']} C"),
            ("Reverse Tm",f"{pair['reverse_tm']} C"),
            ("Amplicon coordinates",f"{pair['amplicon_start_1based']}–{pair['amplicon_end_1based']}"),
        ]
    return rows


def create_pdf(path, result, enzyme, ladder_name, gel_file, map_file):
    styles=getSampleStyleSheet()
    story=[
        Paragraph("PCR-RFLP candidate assay report",styles["Title"]),
        Spacer(1,12),
        Paragraph(escape(DISCLAIMER),styles["BodyText"]),
        Spacer(1,12),
    ]
    details=[["Field","Value"]] + [[escape(str(a)),escape(str(b))] for a,b in _summary_rows(result,enzyme,ladder_name)]
    detail_table=Table(details,colWidths=[140,370],repeatRows=1,hAlign="LEFT")
    detail_table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#dae8f5")),
        ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#cbd5e1")),
        ("FONTSIZE",(0,0),(-1,-1),8),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f7fafc")]),
    ]))
    story.extend([detail_table,Spacer(1,14)])
    if enzyme:
        story.append(Paragraph("Predicted digestion fragments (base pairs)",styles["Heading2"]))
        wt=", ".join(map(str,enzyme["fragments_wt"]))
        alt=", ".join(map(str,enzyme["fragments_mut"]))
        het=", ".join(map(str,sorted(enzyme["fragments_wt"]+enzyme["fragments_mut"],reverse=True)))
        for key,value in [("WT",wt),("Heterozygote",het),("Mutant",alt)]:
            story.append(Paragraph(f"<b>{key}:</b> {escape(value)} bp",styles["BodyText"]))
        story.append(Paragraph("Scoring explanation: "+escape(enzyme.get("design_notes","")),styles["BodyText"]))
    if gel_file:
        story.extend([Spacer(1,10),Paragraph("Predicted virtual gel",styles["Heading2"]),
                      PDFImage(str(gel_file),width=6.7*inch,height=4.0*inch)])
    if map_file:
        story.extend([Spacer(1,8),Paragraph("Primer-binding diagram",styles["Heading2"]),
                      PDFImage(str(map_file),width=6.7*inch,height=1.65*inch)])
    story.extend([Spacer(1,10),Paragraph("Limitations",styles["Heading2"])])
    for warning in result["warnings"]:
        story.append(Paragraph("• "+escape(str(warning)),styles["BodyText"]))
    ladder=LADDERS[ladder_name]
    story.append(Paragraph("Ladder: "+escape(ladder["note"]),styles["BodyText"]))
    if ladder["source"]:
        story.append(Paragraph("Manufacturer source: "+escape(ladder["source"]),styles["BodyText"]))
    SimpleDocTemplate(str(path),pagesize=(612,792),rightMargin=42,leftMargin=42,
                      topMargin=42,bottomMargin=42).build(story)


def create_word(path,result,enzyme,ladder_name,gel_file,map_file):
    doc=Document()
    doc.add_heading("PCR-RFLP candidate assay report",0)
    doc.add_paragraph(DISCLAIMER)
    details=doc.add_table(rows=1,cols=2)
    details.style="Light Shading Accent 1"
    details.rows[0].cells[0].text="Field"
    details.rows[0].cells[1].text="Value"
    for label,value in _summary_rows(result,enzyme,ladder_name):
        a,b=details.add_row().cells
        a.text=str(label);b.text=str(value)
    if enzyme:
        doc.add_heading("Predicted fragments (bp)",level=2)
        for label,values in [
            ("Wild type",enzyme["fragments_wt"]),
            ("Heterozygote",sorted(enzyme["fragments_wt"]+enzyme["fragments_mut"],reverse=True)),
            ("Mutant",enzyme["fragments_mut"]),
        ]:
            doc.add_paragraph(f"{label}: {', '.join(map(str,values))} bp")
        doc.add_paragraph("Enzyme-priority rationale: "+enzyme.get("design_notes",""))
    if gel_file:
        doc.add_heading("Virtual gel — complete-digestion prediction",level=2)
        doc.add_picture(str(gel_file),width=Inches(6.1))
    if map_file:
        doc.add_heading("Primer binding positions",level=2)
        doc.add_picture(str(map_file),width=Inches(6.1))
    doc.add_heading("Limitations and reproducibility",level=2)
    for warning in result["warnings"]:
        doc.add_paragraph(str(warning),style="List Bullet")
    doc.add_paragraph("Selected ladder: "+ladder_name)
    doc.add_paragraph("Manufacturer documentation: "+LADDERS[ladder_name]["source"])
    doc.save(path)


def create_html(path,result,enzyme,ladder_name):
    summary="".join(f"<tr><th>{escape(str(a))}</th><td>{escape(str(b))}</td></tr>"
                    for a,b in _summary_rows(result,enzyme,ladder_name))
    ranks=pd.DataFrame(result["enzymes"]).to_html(index=False,escape=True) if result["enzymes"] else "<p>No enzymes.</p>"
    intro=f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>PCR-RFLP report</title><style>body{{font-family:system-ui,Arial;max-width:1050px;
margin:auto;padding:25px;color:#23344a}}table{{border-collapse:collapse;width:100%;font-size:13px}}
td,th{{border:1px solid #d4dfeb;padding:6px;word-break:break-word}}th{{background:#e8f1fa}}
img{{max-width:100%}}</style></head><body><h1>PCR-RFLP assay report</h1>
<p><b>{escape(DISCLAIMER)}</b></p><h2>Assay summary</h2><table>{summary}</table>
<h2>Predicted enzyme rankings</h2>{ranks}
<h2>Virtual gel</h2><img src="virtual_gel.png" alt="In-silico gel">
<h2>Primer position diagram</h2><img src="primer_map.png" alt="Primer map">"""
    intro += "<h2>Highlighted DNA sequence</h2><iframe src='sequence_highlighted.html' width='100%' height='440' title='Sequence positions'></iframe>"
    intro += "<p>Images and iframe are included in the ZIP; open the HTML after extracting the folder.</p></body></html>"
    path.write_text(intro,encoding="utf-8")


def build_export_bundle(result, input_sequence, ladder_name, selected_enzyme):
    """Build all files in a temp directory; return zip bytes and selected individual bytes."""
    with tempfile.TemporaryDirectory() as d:
        folder=Path(d)
        enzyme=selected_enzyme
        gel_file=None
        map_file=None
        if enzyme and result["selected_primer"]:
            gel_file=folder/"virtual_gel.png"
            render_ladder_gel(result["selected_primer"]["product_size"],enzyme,ladder_name,gel_file)
        if result["selected_primer"]:
            map_file=folder/"primer_map.png"
            render_sequence_figure(result,map_file)
            (folder/"sequence_highlighted.html").write_text(sequence_markup(result),encoding="utf-8")
        name="".join(c if c.isalnum() or c in "_-." else "_" for c in result["name"])[:80] or "variant"
        (folder/"input_reference.fasta").write_text(f">{name}\n"+ "\n".join(
            input_sequence[i:i+70] for i in range(0,len(input_sequence),70))+"\n",encoding="utf-8")
        chosen={"ladder":ladder_name,"enzyme":enzyme["enzyme"] if enzyme else None,
                "ladder_product":LADDERS[ladder_name]}
        (folder/"analysis.json").write_text(json.dumps({"assay":result,"selected":chosen},indent=2),
                                            encoding="utf-8")
        pd.DataFrame(result["primers"]).to_csv(folder/"all_primers.csv",index=False)
        pd.DataFrame(result["enzymes"]).to_csv(folder/"ranked_enzymes.csv",index=False)
        fragment_rows=[]
        if enzyme:
            for genotype, fragments in [
                ("WT",enzyme["fragments_wt"]),
                ("Heterozygous",sorted(enzyme["fragments_wt"]+enzyme["fragments_mut"],reverse=True)),
                ("Mutant",enzyme["fragments_mut"]),
            ]:
                for bp in fragments:
                    fragment_rows.append({"genotype":genotype,"fragment_size_bp":bp})
        pd.DataFrame(fragment_rows,columns=["genotype","fragment_size_bp"]).to_csv(
            folder/"expected_fragments.csv",index=False)
        create_pdf(folder/"assay_report.pdf",result,enzyme,ladder_name,gel_file,map_file)
        create_word(folder/"assay_report.docx",result,enzyme,ladder_name,gel_file,map_file)
        create_html(folder/"assay_report.html",result,enzyme,ladder_name)
        notes="\n".join([
            "PCR-RFLP research-use bundle",
            DISCLAIMER,
            "1. Open assay_report.pdf or assay_report.docx for the summary.",
            "2. Open assay_report.html after extracting the whole ZIP.",
            "3. Review ranked_enzymes.csv; grade is in-silico priority, NOT enzyme activity.",
            "4. Primer sequence highlights are in sequence_highlighted.html.",
            "5. Gel assumes complete digestion and does not predict real band intensity.",
            "6. Position is relative to input_reference.fasta, not necessarily HGVS.",
            "7. Vendor ladder reference: "+LADDERS[ladder_name]["source"],
            "8. Keep experimental records; confirm primers and enzymes before ordering.",
        ])
        (folder/"README_RESULTS.txt").write_text(notes,encoding="utf-8")
        individual={}
        for filename in ["assay_report.pdf","assay_report.docx","virtual_gel.png",
                         "primer_map.png","ranked_enzymes.csv","sequence_highlighted.html"]:
            path=folder/filename
            if path.exists():
                individual[filename]=path.read_bytes()
        out=io.BytesIO()
        with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as archive:
            for item in sorted(folder.iterdir()):
                archive.write(item,item.name)
        return out.getvalue(), individual
