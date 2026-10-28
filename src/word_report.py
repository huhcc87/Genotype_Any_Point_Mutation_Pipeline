from docx import Document
from docx.shared import Inches
from pathlib import Path

def generate_word_report(
    name: str,
    mutation_info: dict,
    enzymes: list,
    primers: list,
    gel_path: str = None,
    output_path: str = "report.docx"
):
    doc = Document()
    doc.add_heading(f"PCR-RFLP Design Report: {name}", 0)
    doc.add_paragraph(f"Date: {mutation_info.get('date', '')}")

    doc.add_heading("Mutation", level=1)
    doc.add_paragraph(f"Position (1-based): {mutation_info.get('position', 'N/A')}")
    doc.add_paragraph(f"Change: {mutation_info.get('ref', '?')} → {mutation_info.get('alt', '?')}")
    doc.add_paragraph(f"Sequence length: {mutation_info.get('seq_length', 'N/A')} bp")

    if enzymes:
        best = enzymes[0]
        doc.add_heading("Recommended Enzyme", level=1)
        doc.add_paragraph(f"{best['enzyme']} (site: {best['site']})")

    doc.add_heading("Primers", level=1)
    for p in primers[:3]:
        doc.add_paragraph(
            f"Pair {p['pair']} ({p.get('type', 'natural')}): "
            f"{p['forward']} / {p['reverse']} "
            f"(Tm {p['forward_tm']}°C / {p['reverse_tm']}°C, {p['product_size']} bp)"
        )

    if gel_path and Path(gel_path).exists():
        doc.add_heading("Virtual Gel", level=1)
        doc.add_picture(gel_path, width=Inches(5))

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
